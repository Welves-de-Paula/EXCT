import tkinter as tk
from tkinter import ttk
import os
from openpyxl import load_workbook

def get_latest_file_from_origin():
    origin_dir = os.path.join(os.path.dirname(__file__), "..", "data", "origin")
    if not os.path.isdir(origin_dir):
        return None
    files = [os.path.join(origin_dir, f) for f in os.listdir(origin_dir) if os.path.isfile(os.path.join(origin_dir, f))]
    if not files:
        return None
    return max(files, key=os.path.getmtime)

def main(filepath=None, title="Layout com Botão Validar"):
    # Se filepath não for passado, tenta pegar o arquivo mais recente de data/origin
    if filepath is None:
        filepath = get_latest_file_from_origin()

    file_content = ""
    table_data = []
    table_headers = []
    file_is_compatible = False  # Flag para compatibilidade

    # Variável para armazenar o resultado do identifier
    identifier_result = None

    if filepath and os.path.isfile(filepath):
        ext = os.path.splitext(filepath)[1].lower()
        if ext == ".xlsx":
            try:
                wb = load_workbook(filepath, read_only=True, data_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                if rows:
                    table_headers = [str(h) if h is not None else "" for h in rows[0]]
                    table_data = [
                        [str(cell) if cell is not None else "" for cell in row]
                        for row in rows[1:]
                    ]
                    file_is_compatible = True
                    # Corrige o caminho do identifier.py para ../script/identifier.py
                    import importlib.util
                    import sys
                    identifier_path = os.path.join(os.path.dirname(__file__), "..", "script", "identifier.py")
                    spec = importlib.util.spec_from_file_location("identifier", identifier_path)
                    identifier = importlib.util.module_from_spec(spec)
                    sys.modules["identifier"] = identifier
                    spec.loader.exec_module(identifier)
                    # Supondo que identifier.py tenha uma função 'main' que recebe headers e dados
                    if hasattr(identifier, "main"):
                        identifier_result = identifier.main(table_headers, table_data)
            except Exception as e:
                file_content = f"Erro ao ler arquivo Excel: {e}"
        else:
            file_content = "Arquivo não suportado. Por favor, utilize apenas arquivos .xlsx."

    # Se o arquivo não for compatível, abre o incompatible.py imediatamente
    if not file_is_compatible:
        import importlib.util
        incompatible_path = os.path.join(os.path.dirname(__file__), "incompatible.py")
        spec = importlib.util.spec_from_file_location("incompatible", incompatible_path)
        incompatible = importlib.util.module_from_spec(spec)
        import sys
        sys.modules["incompatible"] = incompatible
        spec.loader.exec_module(incompatible)
        if hasattr(incompatible, "main"):
            incompatible.main("Arquivo incompatível. Por favor, utilize apenas arquivos .xlsx válidos.")
        return  # Não segue para a interface principal

    def on_validar():
        print("Botão Validar clicado!")
        if filepath:
            print(f"Arquivo recebido: {filepath}")
        if file_is_compatible:
            # Corrige o caminho do identifier.py para ../script/identifier.py
            import importlib.util
            import sys
            identifier_path = os.path.join(os.path.dirname(__file__), "..", "script", "identifier.py")
            spec = importlib.util.spec_from_file_location("identifier", identifier_path)
            identifier = importlib.util.module_from_spec(spec)
            sys.modules["identifier"] = identifier
            spec.loader.exec_module(identifier)
            # Supondo que identifier.py tenha uma função 'main' que recebe headers e dados
            if hasattr(identifier, "main"):
                identifier.main(table_headers, table_data)
        else:
            # Abre incompatible.py com mensagem
            import importlib.util
            incompatible_path = os.path.join(os.path.dirname(__file__), "incompatible.py")
            spec = importlib.util.spec_from_file_location("incompatible", incompatible_path)
            incompatible = importlib.util.module_from_spec(spec)
            sys.modules["incompatible"] = incompatible
            spec.loader.exec_module(incompatible)
            # Supondo que incompatible.py tenha uma função 'main' que recebe uma mensagem
            if hasattr(incompatible, "main"):
                incompatible.main("Arquivo incompatível. Por favor, utilize apenas arquivos .xlsx válidos.")

    root = tk.Tk()
    root.title(title)
    root.state('zoomed')  # Maximiza a janela

    frame = tk.Frame(root, padx=20, pady=20)
    frame.pack()

    btn_validar = tk.Button(frame, text="Validar", command=on_validar)
    btn_validar.pack(side="left")

    # Exibe apenas mensagem de erro se houver, não exibe tabela
    display_frame = tk.Frame(root, padx=20, pady=20)
    display_frame.pack(fill="both", expand=True)

    if file_content:
        text_widget = tk.Text(display_frame, wrap="none")
        text_widget.pack(fill="both", expand=True)
        text_widget.insert("1.0", file_content)
        text_widget.config(state="disabled")

    # Exemplo seguro de uso do after:
    def safe_after():
        if root.winfo_exists():
            # ...código que precisa ser executado periodicamente...
            root.after(100, safe_after)
    # safe_after()  # Descomente se precisar usar after

    root.mainloop()